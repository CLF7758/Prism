"""T3：文档导出必须保真，而且不能误导用户。

导出是「用户把这个文件拿走」的那一步，所以三件事要钉住：

* HTML 里如果留着 ``prism://attach/N``，浏览器打开就是一片裂图 —— 用户
  以为导出了能看的网页，其实什么也没看到；
* DOCX 里的图片如果被 Qt 重新编码过，画质就再也回不来了（用户硬要求：
  导入什么导出就该是什么）；
* 对话框里出现的格式必须真的能导出，否则就是在误导用户。

PDF 走纯 Qt 的 QTextDocument + QPdfWriter，所以这里的测试不需要
QtWebEngine。
"""
import base64
import io
import os
import re
import zipfile

import pytest
from PyQt6 import QtCore, QtGui

from prism.fileio import document_export
from prism.fileio.attachments import attachment_url

#: 一张 3x3 的红色 PNG，和 conftest 里那张同一个来源。
def _png_bytes():
    image = QtGui.QImage(3, 3, QtGui.QImage.Format.Format_RGB32)
    image.fill(QtGui.QColor('red'))
    buffer = QtCore.QBuffer()
    buffer.open(QtCore.QIODevice.OpenModeFlag.WriteOnly)
    image.save(buffer, 'PNG')
    return bytes(buffer.data())


def _photo_bytes(quality=95):
    """A real JPEG, so a re-encode would show up as different bytes."""
    from PIL import Image

    buffer = io.BytesIO()
    Image.new('RGB', (64, 48), (200, 30, 30)).save(
        buffer, 'JPEG', quality=quality)
    return buffer.getvalue()


def _attachment(blob, name):
    mime = 'image/jpeg' if name.lower().endswith(('.jpg', '.jpeg')) \
        else 'image/png'
    return (name, mime, blob)


# ── The format list ──────────────────────────────────────────────


def test_every_offered_format_has_a_writer():
    """任务书 T3：「导出为 HTML、DOCX 或 PDF 时不误导用户」。

    能出现在对话框里的格式，必须真的有实现。
    """
    for entry in document_export.EXPORT_FORMATS:
        assert entry.extension.startswith('.')
        assert entry.label
        assert callable(entry.writer), entry


def test_pdf_is_offered():
    extensions = {entry.extension for entry in document_export.EXPORT_FORMATS}
    assert '.pdf' in extensions
    assert '.html' in extensions


def test_the_chosen_format_follows_the_path():
    """用户在对话框里打了 note.pdf，就不该存成 note.pdf.docx。"""
    entry = document_export.writer_for('C:/tmp/note.pdf')
    assert entry.extension == '.pdf'
    assert document_export.with_extension('C:/tmp/note.pdf', entry) == \
        'C:/tmp/note.pdf'


def test_the_chosen_format_follows_the_filter_when_there_is_no_extension():
    entry = document_export.writer_for('C:/tmp/note', 'PDF 文档 (*.pdf)')
    assert entry.extension == '.pdf'
    assert document_export.with_extension('C:/tmp/note', entry) == \
        'C:/tmp/note.pdf'


def test_a_web_page_keeps_its_htm_spelling():
    entry = document_export.writer_for('C:/tmp/note.htm')
    assert entry.extension == '.html'
    assert document_export.with_extension('C:/tmp/note.htm', entry) == \
        'C:/tmp/note.htm'


def test_an_unknown_extension_falls_back_to_word():
    entry = document_export.writer_for('C:/tmp/note.doc', 'Word Document (*.docx)')
    assert entry.extension == '.docx'
    assert document_export.with_extension('C:/tmp/note.doc', entry) == \
        'C:/tmp/note.doc.docx'


# ── HTML export ──────────────────────────────────────────────────


def test_exported_html_carries_no_prism_urls():
    blob = _png_bytes()
    html = f'<p>正文</p><p><img src="{attachment_url(4)}"></p>'
    exported = document_export.html_with_embedded_images(
        html, {4: _attachment(blob, 'shot.png')})

    assert 'prism://' not in exported
    assert 'data:image/png;base64,' in exported
    assert '正文' in exported


def test_exported_html_embeds_the_stored_bytes_unchanged():
    """内嵌的必须是原字节：导出不能是第二个降质量的地方。"""
    blob = _photo_bytes()
    html = f'<img src="{attachment_url(9)}">'
    exported = document_export.html_with_embedded_images(
        html, {9: _attachment(blob, 'photo.jpg')})

    payload = re.search(r'base64,([^"]+)', exported).group(1)
    assert base64.b64decode(payload) == blob
    assert 'data:image/jpeg;base64,' in exported


def test_exported_html_leaves_a_picture_it_cannot_find_alone():
    """缺附件时保留原样，不抛异常，也不静默删掉整段文字。"""
    html = f'<p>留着的字</p><img src="{attachment_url(99)}">'
    exported = document_export.html_with_embedded_images(html, {})
    assert '留着的字' in exported
    assert 'prism://attach/99' in exported


def test_exported_html_absorbed_inline_data_images():
    """页面可能把还没入库的图片以 data: URL 送回来，导出时不该漏掉。"""
    blob = _png_bytes()
    payload = base64.b64encode(blob).decode('ascii')
    html = f'<img src="data:image/png;base64,{payload}">'
    exported = document_export.html_with_embedded_images(html, {})
    assert exported == html


def test_html_export_writes_a_file(tmp_path):
    blob = _png_bytes()
    target = tmp_path / 'note.html'
    count = document_export.export_note_to_html(
        f'<p>标题</p><img src="{attachment_url(1)}">',
        {1: _attachment(blob, 'a.png')}, str(target))

    assert count == 1
    text = target.read_text(encoding='utf-8')
    assert '标题' in text
    assert 'prism://' not in text
    # 一个能直接双击打开的文件，而不是一段 HTML 片段。
    assert '<html' in text.lower()
    assert 'charset' in text.lower()


def test_html_export_does_not_touch_the_source():
    attachments = {1: _attachment(_png_bytes(), 'a.png')}
    before = dict(attachments)
    html = f'<img src="{attachment_url(1)}">'
    document_export.html_with_embedded_images(html, attachments)
    assert attachments == before


# ── PDF export ───────────────────────────────────────────────────


def test_pdf_export_writes_a_pdf(qapp, tmp_path):
    target = tmp_path / 'note.pdf'
    pages = document_export.export_note_to_pdf(
        '<h1>标题</h1><p>正文</p>', {}, str(target))

    assert target.exists()
    with open(target, 'rb') as handle:
        assert handle.read(5) == b'%PDF-'
    assert pages >= 1


def test_pdf_export_resolves_board_pictures(qapp, tmp_path, view):
    """图片必须真的画进去，而不是留一个空位。

    和一份只有文字的 PDF 比：图片会多出一个 image XObject，文件也要大
    一截。用相对比较而不是写死的阈值，Qt 换版本时不会假报警。
    """
    view.scene.attachments.clear()
    view.scene.attachments[3] = _attachment(_photo_bytes(), 'photo.jpg')
    with_picture = tmp_path / 'with.pdf'
    document_export.export_note_to_pdf(
        f'<p><img src="{attachment_url(3)}"></p>',
        view.scene.attachments, str(with_picture))

    plain = tmp_path / 'plain.pdf'
    document_export.export_note_to_pdf(
        '<p>just text</p>', {}, str(plain))

    drawn = with_picture.read_bytes()
    text_only = plain.read_bytes()
    assert drawn.startswith(b'%PDF-')
    assert len(drawn) > len(text_only)
    assert drawn.count(b'/XObject') > text_only.count(b'/XObject')


def test_pdf_export_limits_a_picture_to_the_page_width():
    html = '<img src="prism://attach/1">'
    fitted = document_export.html_with_fitted_images(
        html, document_export.MAX_PAGE_IMAGE_WIDTH)
    assert f'width="{document_export.MAX_PAGE_IMAGE_WIDTH}"' in fitted


def test_pdf_export_keeps_an_explicit_picture_size():
    html = '<img src="prism://attach/1" width="120" height="90">'
    fitted = document_export.html_with_fitted_images(
        html, document_export.MAX_PAGE_IMAGE_WIDTH)
    assert 'width="120"' in fitted
    assert f'width="{document_export.MAX_PAGE_IMAGE_WIDTH}"' not in fitted


def test_pdf_export_handles_a_self_closing_tag():
    fitted = document_export.html_with_fitted_images(
        '<img src="prism://attach/1" />',
        document_export.MAX_PAGE_IMAGE_WIDTH)
    assert fitted == ('<img src="prism://attach/1" '
                      f'width="{document_export.MAX_PAGE_IMAGE_WIDTH}">')


def test_pdf_export_handles_a_style_sized_picture():
    html = '<img src="prism://attach/1" style="width: 200px">'
    assert document_export.html_with_fitted_images(
        html, document_export.MAX_PAGE_IMAGE_WIDTH) == html


def test_pdf_export_reports_a_target_it_cannot_write(qapp, tmp_path):
    with pytest.raises(document_export.DocumentExportError):
        document_export.export_note_to_pdf(
            '<p>x</p>', {}, str(tmp_path / 'missing' / 'note.pdf'))


# ── DOCX export keeps the bytes ──────────────────────────────────


def test_docx_export_embeds_the_original_picture_bytes(qapp, view, tmp_path):
    """导入什么导出就该是什么：JPEG 不能因为导出被重新编码一次。"""
    blob = _photo_bytes()
    view.scene.attachments.clear()
    view.scene.attachments[5] = _attachment(blob, 'photo.jpg')
    target = tmp_path / 'note.docx'

    from prism.fileio.word import export_note_to_docx
    count = export_note_to_docx(
        f'<p><img src="{attachment_url(5)}"></p>',
        view.scene.attachments, str(target))

    assert count == 1
    with zipfile.ZipFile(target) as archive:
        media = [name for name in archive.namelist()
                 if name.startswith('word/media/')]
        assert len(media) == 1
        assert archive.read(media[0]) == blob


def test_docx_export_still_converts_a_format_word_cannot_embed(qapp, view,
                                                              tmp_path):
    """Word 收不了 WebP，那就转成 PNG —— 而不是把它整张丢掉。"""
    from PIL import Image

    buffer = io.BytesIO()
    Image.new('RGB', (32, 24), 'blue').save(buffer, 'WEBP')
    view.scene.attachments.clear()
    view.scene.attachments[6] = ('shot.webp', 'image/webp', buffer.getvalue())
    target = tmp_path / 'note.docx'

    from prism.fileio.word import export_note_to_docx
    count = export_note_to_docx(
        f'<p><img src="{attachment_url(6)}"></p>',
        view.scene.attachments, str(target))

    assert count == 1
    with zipfile.ZipFile(target) as archive:
        media = [name for name in archive.namelist()
                 if name.startswith('word/media/')]
        assert media
        assert archive.read(media[0]).startswith(b'\x89PNG')


def test_docx_export_passes_a_tiff_through_untouched(qapp, view, tmp_path):
    """TIFF 也在 Word 能收的清单里，所以同样不该被重编码。"""
    from PIL import Image

    buffer = io.BytesIO()
    Image.new('RGB', (32, 24), 'blue').save(buffer, 'TIFF')
    blob = buffer.getvalue()
    view.scene.attachments.clear()
    view.scene.attachments[11] = ('scan.tif', 'image/tiff', blob)
    target = tmp_path / 'note.docx'

    from prism.fileio.word import export_note_to_docx
    export_note_to_docx(
        f'<p><img src="{attachment_url(11)}"></p>',
        view.scene.attachments, str(target))

    with zipfile.ZipFile(target) as archive:
        media = [name for name in archive.namelist()
                 if name.startswith('word/media/')]
        assert media
        assert archive.read(media[0]) == blob


def test_an_unreadable_attachment_is_skipped_not_fatal(qapp, view, tmp_path):
    view.scene.attachments.clear()
    view.scene.attachments[8] = ('broken.png', 'image/png', b'not an image')
    target = tmp_path / 'note.docx'

    from prism.fileio.word import export_note_to_docx
    export_note_to_docx(
        f'<p>文字</p><p><img src="{attachment_url(8)}"></p>',
        view.scene.attachments, str(target))

    assert target.exists()
    assert os.path.getsize(target) > 0


# ── The whole chain ──────────────────────────────────────────────


def test_a_picture_survives_import_and_export_byte_for_byte(qapp, view,
                                                            tmp_path):
    """端到端：.docx 里的图 → 导入 → 导出，字节必须一模一样。

    这是「导入什么导出就该是什么」的完整证据链。中间要过附件表、编辑器
    的 HTML、python-docx 打包三道手，任何一步偷偷重编码都会在这里露出来。
    """
    import docx as docx_module

    from prism.fileio.import_docx import docx_to_html
    from prism.fileio.word import export_note_to_docx

    blob = _photo_bytes()
    picture = tmp_path / 'photo.jpg'
    picture.write_bytes(blob)
    document = docx_module.Document()
    document.add_paragraph('带图的一段')
    document.add_picture(str(picture))
    source = tmp_path / 'source.docx'
    document.save(str(source))

    view.scene.attachments.clear()

    def saver(data, extension):
        identifier = max(view.scene.attachments.keys(), default=0) + 1
        view.scene.attachments[identifier] = (
            f'imported.{extension}', f'image/{extension}', data)
        return attachment_url(identifier)

    markup = docx_to_html(str(source), attachment_saver=saver)

    assert 'prism://attach/1' in markup
    assert view.scene.attachments[1][2] == blob        # 导入没有重编码

    target = tmp_path / 'roundtrip.docx'
    export_note_to_docx(markup, view.scene.attachments, str(target))

    with zipfile.ZipFile(target) as archive:
        media = [name for name in archive.namelist()
                 if name.startswith('word/media/')]
        assert len(media) == 1
        assert archive.read(media[0]) == blob           # 导出也没有
