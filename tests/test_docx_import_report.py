"""T3：Word 导入能带过来什么、带不过来什么，必须说得出来。

任务书要求「不支持的特性必须明确提示，而不是静默丢失」。这一条以前完全
没有：页眉、脚注、批注、修订、公式、文本框都是悄悄不见的，用户以为整份
文件都进来了，然后拿着一份不是自己写的文档去做事。

这些测试用 python-docx 造出真的 .docx，所以覆盖的是用户文件里真实的
WordprocessingML，不是构造出来的字符串。
"""
import pytest

docx = pytest.importorskip('docx')

from docx.oxml import OxmlElement, parse_xml                 # noqa: E402
from docx.oxml.ns import nsdecls, qn                         # noqa: E402

from prism.fileio.import_docx import (                       # noqa: E402
    ImportReport, SKIPPED_FEATURES, docx_to_html,
    scan_skipped_features)


def save_document(document, tmp_path, name='note.docx'):
    path = tmp_path / name
    document.save(str(path))
    return str(path)


def _append(paragraph, xml):
    paragraph._p.append(parse_xml(xml))


def _document_with_header(text='页眉内容'):
    document = docx.Document()
    document.add_paragraph('正文')
    document.sections[0].header.paragraphs[0].text = text
    return document


# ── A plain document says nothing ────────────────────────────────


def test_a_plain_document_reports_nothing(tmp_path):
    document = docx.Document()
    document.add_paragraph('就是一段普通的文字')
    report = ImportReport()

    html = docx_to_html(save_document(document, tmp_path), report=report)

    assert '普通的文字' in html
    assert report.skipped == {}
    assert report.summary() is None


def test_the_report_counts_the_pictures(tmp_path, imgfilename3x3):
    from docx.shared import Inches

    document = docx.Document()
    document.add_picture(imgfilename3x3, width=Inches(2))
    report = ImportReport()

    def saver(data, extension):
        return 'prism://attach/1'

    docx_to_html(save_document(document, tmp_path), attachment_saver=saver,
                 report=report)
    assert report.images == 1
    assert report.skipped == {}


def test_importing_without_a_report_still_works(tmp_path):
    """老调用点不传 report 也不能出问题。"""
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = '页眉'
    assert docx_to_html(save_document(document, tmp_path)).strip()


# ── Each feature is noticed ──────────────────────────────────────


def test_headers_are_reported(tmp_path):
    report = ImportReport()
    docx_to_html(save_document(_document_with_header(), tmp_path),
                 report=report)
    assert 'header' in report.skipped


def test_a_header_is_really_not_in_the_html(tmp_path):
    """报告的东西必须是真的没带过来，不能只是报着玩。"""
    html = docx_to_html(save_document(_document_with_header('页眉内容'),
                                      tmp_path))
    assert '页眉内容' not in html


def test_footnotes_are_reported(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph('正文')
    run = OxmlElement('w:r')
    reference = OxmlElement('w:footnoteReference')
    reference.set(qn('w:id'), '2')
    run.append(reference)
    paragraph._p.append(run)

    report = ImportReport()
    docx_to_html(save_document(document, tmp_path), report=report)
    assert 'footnote' in report.skipped


def test_comments_are_reported(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph('正文')
    run = OxmlElement('w:r')
    reference = OxmlElement('w:commentReference')
    reference.set(qn('w:id'), '1')
    run.append(reference)
    paragraph._p.append(run)

    report = ImportReport()
    docx_to_html(save_document(document, tmp_path), report=report)
    assert 'comment' in report.skipped


def test_tracked_changes_are_reported(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph('正文')
    _append(paragraph, (
        f'<w:ins {nsdecls("w")} w:id="1" w:author="a" '
        'w:date="2024-01-01T00:00:00Z">'
        '<w:r><w:t>新插进来的一句</w:t></w:r></w:ins>'))

    report = ImportReport()
    docx_to_html(save_document(document, tmp_path), report=report)
    assert 'revision' in report.skipped


def test_equations_are_reported(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph('公式：')
    paragraph._p.append(parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/'
        'officeDocument/2006/math"><m:r><m:t>x=1</m:t></m:r></m:oMath>'))

    report = ImportReport()
    docx_to_html(save_document(document, tmp_path), report=report)
    assert 'formula' in report.skipped


def test_text_boxes_are_reported(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph('正文')
    _append(paragraph, (
        f'<w:r {nsdecls("w")}><w:txbxContent>'
        '<w:p><w:r><w:t>框里的话</w:t></w:r></w:p>'
        '</w:txbxContent></w:r>'))

    report = ImportReport()
    docx_to_html(save_document(document, tmp_path), report=report)
    assert 'textbox' in report.skipped


def test_several_features_are_all_reported(tmp_path):
    document = _document_with_header()
    paragraph = document.add_paragraph('正文')
    _append(paragraph, (
        f'<w:ins {nsdecls("w")} w:id="1" w:author="a" '
        'w:date="2024-01-01T00:00:00Z">'
        '<w:r><w:t>改过的</w:t></w:r></w:ins>'))

    report = ImportReport()
    docx_to_html(save_document(document, tmp_path), report=report)
    assert set(report.skipped) == {'header', 'revision'}


def test_scanning_an_empty_body_reports_nothing():
    document = docx.Document()
    report = ImportReport()
    scan_skipped_features(document.element.body, report)
    assert report.skipped == {}


# ── What the user is told ────────────────────────────────────────


def test_the_summary_names_the_features_in_words(tmp_path):
    report = ImportReport()
    docx_to_html(save_document(_document_with_header(), tmp_path),
                 report=report)

    summary = report.summary()
    assert summary
    assert 'headers and footers' in summary
    # 界面上是中文，catalogue 里必须有对应的说法。
    from prism.i18n.zh_CN import zh_CN
    for feature in report.skipped:
        assert SKIPPED_FEATURES[feature] in zh_CN, feature


def test_the_summary_counts_scaled_pictures(tmp_path, monkeypatch):
    from PIL import Image
    import io

    from docx.shared import Inches
    from prism.fileio import import_docx

    document = docx.Document()
    buffer = io.BytesIO()
    Image.new('RGB', (400, 200), 'red').save(buffer, 'PNG')
    picture = tmp_path / 'big.png'
    picture.write_bytes(buffer.getvalue())
    document.add_picture(str(picture), width=Inches(2))
    document.sections[0].header.paragraphs[0].text = '页眉'

    monkeypatch.setattr(import_docx, 'MAX_IMAGE_PIXELS', 4000)
    report = ImportReport()
    docx_to_html(save_document(document, tmp_path), attachment_saver=_noop,
                 report=report)

    assert report.images_scaled == 1
    assert '1' in report.summary()


def _noop(data, extension):
    return 'prism://attach/1'


def test_the_report_can_be_logged_as_data(tmp_path):
    report = ImportReport()
    docx_to_html(save_document(_document_with_header(), tmp_path),
                 report=report)
    as_dict = report.as_dict()
    assert as_dict['skipped'] == dict(report.skipped)
    assert isinstance(as_dict['images'], int)
    assert isinstance(as_dict['images_scaled'], int)
