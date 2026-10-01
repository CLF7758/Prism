"""The .docx reader maps Word's formatting onto what the note editor shows.

The reader is handed real .docx files built here with python-docx, so the
tests exercise the same WordprocessingML a user's document carries.
"""
import io

import pytest

docx = pytest.importorskip('docx')

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT         # noqa: E402
from docx.oxml import OxmlElement                              # noqa: E402
from docx.oxml.ns import qn                                    # noqa: E402
from docx.shared import Inches, Pt, RGBColor                   # noqa: E402

from prism.fileio import import_docx                           # noqa: E402
from prism.fileio.import_base import ImportFileError           # noqa: E402
from prism.fileio.import_docx import docx_to_html              # noqa: E402


class Saver:
    """Stand-in for the board, remembering what it was asked to store."""

    def __init__(self):
        self.images = {}

    def __call__(self, data, extension):
        identifier = len(self.images) + 1
        self.images[identifier] = (data, extension)
        return f'prism://attach/{identifier}'

    def total_bytes(self):
        return sum(len(data) for data, _extension in self.images.values())


@pytest.fixture
def saver():
    return Saver()


def save_document(document, tmp_path, name='note.docx'):
    path = tmp_path / name
    document.save(str(path))
    return str(path)


def add_list_level(paragraph, level, num_id='1'):
    """Give a paragraph an explicit list level, the way Word writes one."""
    properties = paragraph._p.get_or_add_pPr()
    numbering = OxmlElement('w:numPr')
    item = OxmlElement('w:ilvl')
    item.set(qn('w:val'), str(level))
    numbering.append(item)
    item = OxmlElement('w:numId')
    item.set(qn('w:val'), str(num_id))
    numbering.append(item)
    properties.append(numbering)


def add_hyperlink(paragraph, url, text):
    """Append a hyperlink to a paragraph, the way Word writes one."""
    part = paragraph.part
    relation_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), relation_id)
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    style = OxmlElement('w:rStyle')
    style.set(qn('w:val'), 'Hyperlink')
    properties.append(style)
    run.append(properties)
    node = OxmlElement('w:t')
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


# ── Character formatting ──────────────────────────────────────────────

def test_run_formatting_survives_the_import(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run('styled')
    run.bold = True
    run.italic = True
    run.underline = True
    run.font.strike = True
    run.font.size = Pt(20)
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(0x11, 0x22, 0x33)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    html = docx_to_html(save_document(document, tmp_path))

    for expected in ('<b>', '<i>', '<u>', '<s>', 'font-size:20pt',
                     'font-family:\'Arial\'', 'color:#112233',
                     'background-color:#ffff00'):
        assert expected in html


def test_superscript_and_subscript(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run('2')
    run.font.superscript = True
    paragraph.add_run(' and ')
    run = paragraph.add_run('3')
    run.font.subscript = True
    html = docx_to_html(save_document(document, tmp_path))
    assert '<sup>' in html and '<sub>' in html


def test_hidden_text_is_left_out(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('shown ')
    run = paragraph.add_run('secret')
    run.font.hidden = True
    html = docx_to_html(save_document(document, tmp_path))
    assert 'shown' in html
    assert 'secret' not in html


def test_tabs_and_breaks(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph('one')
    run = paragraph.add_run()
    run.add_tab()
    run.add_break()
    paragraph.add_run('two')
    html = docx_to_html(save_document(document, tmp_path))
    assert '&#9;' in html
    assert '<br />' in html


# ── Paragraph formatting ──────────────────────────────────────────────

def test_headings_and_alignment(tmp_path):
    document = docx.Document()
    document.add_heading('Chapter', level=1)
    document.add_heading('Detail', level=4)
    centered = document.add_paragraph('middle')
    centered.alignment = WD_ALIGN_PARAGRAPH.CENTER
    justified = document.add_paragraph('stretched')
    justified.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    html = docx_to_html(save_document(document, tmp_path))

    assert '<h1' in html and 'Chapter' in html
    assert '<h4' in html and 'Detail' in html
    # Qt ignores CSS justification but honours the attribute, so both
    # spellings have to be there.
    assert 'align="center"' in html
    assert 'align="justify"' in html
    assert 'text-align:justify' in html


def test_localised_heading_names(tmp_path):
    document = docx.Document()
    style = document.styles.add_style('标题 2',
                                      docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    paragraph = document.add_paragraph('中文标题')
    paragraph.style = style
    html = docx_to_html(save_document(document, tmp_path))
    assert '<h2' in html


def test_indent_spacing_and_line_height(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph('indented')
    format_ = paragraph.paragraph_format
    format_.left_indent = Inches(0.5)
    format_.first_line_indent = Inches(-0.25)
    format_.line_spacing = 1.5
    format_.space_before = Pt(6)
    format_.space_after = Pt(0)
    html = docx_to_html(save_document(document, tmp_path))

    assert 'margin-left:48px' in html
    assert 'text-indent:-24px' in html
    assert 'line-height:1.5' in html
    assert 'margin-top:8px' in html
    assert 'margin-bottom:0px' in html


def test_empty_paragraphs_are_kept(tmp_path):
    document = docx.Document()
    document.add_paragraph('before')
    document.add_paragraph('')
    document.add_paragraph('after')
    html = docx_to_html(save_document(document, tmp_path))
    assert html.count('<br />') == 1


# ── Lists ─────────────────────────────────────────────────────────────

def test_bullet_and_numbered_styles_become_lists(tmp_path):
    document = docx.Document()
    document.add_paragraph('first', style='List Bullet')
    document.add_paragraph('second', style='List Bullet')
    document.add_paragraph('third', style='List Number')
    html = docx_to_html(save_document(document, tmp_path))

    assert '<ul>' in html and '<ol' in html
    assert html.count('<li>') == 3
    # Two bullets stay in one list instead of one list per item.
    assert html.count('<ul>') == 1


def test_numbering_switched_off_is_not_a_list(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph('plain', style='List Bullet')
    add_list_level(paragraph, 0, num_id='0')
    html = docx_to_html(save_document(document, tmp_path))
    assert '<li>' not in html


def test_nested_list_items(tmp_path):
    writer = import_docx._ListWriter()
    markup = writer.item(0, 'ul', None, '1', 'top')
    markup += writer.item(1, 'ul', None, '2', 'child')
    markup += writer.item(1, 'ul', None, '3', 'child two')
    markup += writer.item(0, 'ul', None, '4', 'top two')
    markup += writer.close()
    assert markup == ('<ul><li>top<ul><li>child</li><li>child two</li></ul>'
                      '</li><li>top two</li></ul>')


def test_numbered_lists_restart_when_the_id_changes(tmp_path):
    writer = import_docx._ListWriter()
    markup = writer.item(0, 'ol', '1', '1', 'one')
    markup += writer.item(0, 'ol', '1', '2', 'one again')
    markup += writer.close()
    assert markup == ('<ol type="1"><li>one</li></ol>'
                      '<ol type="1"><li>one again</li></ol>')


# ── Images, tables, links ─────────────────────────────────────────────

def test_images_are_stored_once_and_keep_their_bytes(tmp_path, saver,
                                                     imgfilename3x3):
    document = docx.Document()
    document.add_picture(imgfilename3x3, width=Inches(2))
    document.add_picture(imgfilename3x3, width=Inches(2))
    html = docx_to_html(save_document(document, tmp_path), saver)

    assert html.count('<img ') == 2
    assert len(saver.images) == 1              # Word reuses the media part
    stored, extension = saver.images[1]
    with open(imgfilename3x3, 'rb') as handle:
        assert stored == handle.read()         # bytes pass through as they are
    assert extension == 'png'
    assert 'width="192"' in html


def test_an_image_too_large_to_show_is_scaled_down(tmp_path, saver,
                                                   monkeypatch):
    from PIL import Image

    document = docx.Document()
    buffer = io.BytesIO()
    Image.new('RGB', (400, 200), 'red').save(buffer, 'PNG')
    picture = tmp_path / 'big.png'
    picture.write_bytes(buffer.getvalue())
    document.add_picture(str(picture))

    # Pretend the editor cannot cope with more than 4000 pixels, so the
    # picture has to shrink even though it fits in the file.
    monkeypatch.setattr(import_docx, 'MAX_IMAGE_PIXELS', 4000)
    docx_to_html(save_document(document, tmp_path), saver)

    stored, extension = saver.images[1]
    with Image.open(io.BytesIO(stored)) as image:
        assert image.size[0] * image.size[1] <= 4000
        assert image.size[0] > image.size[1]
    assert extension == 'png'


def test_table_with_a_merged_cell(tmp_path):
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'heading'
    table.cell(0, 1).text = 'ignored'
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = 'left'
    table.cell(1, 1).text = 'right'
    html = docx_to_html(save_document(document, tmp_path))

    assert '<table' in html
    assert 'colspan="2"' in html
    assert 'heading' in html and 'left' in html and 'right' in html
    # The first row is one merged cell, not two: the cell Word reports for
    # the second grid column is the same one over again.
    rows = html.split('<tr>')
    assert len(rows) == 3
    assert rows[1].count('<td') == 1
    assert rows[2].count('<td') == 2


def test_hyperlinks_are_kept(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph('see ')
    add_hyperlink(paragraph, 'https://example.com/note', 'the note')
    html = docx_to_html(save_document(document, tmp_path))
    assert '<a href="https://example.com/note">' in html
    assert 'the note' in html


# ── Failure paths ─────────────────────────────────────────────────────

def test_a_file_that_is_not_a_docx_reports_an_error(tmp_path):
    broken = tmp_path / 'broken.docx'
    broken.write_text('this is not a Word file', encoding='utf-8')
    with pytest.raises(ImportFileError):
        docx_to_html(str(broken))


def test_an_empty_document_still_yields_markup(tmp_path):
    html = docx_to_html(save_document(docx.Document(), tmp_path))
    assert html.strip()
